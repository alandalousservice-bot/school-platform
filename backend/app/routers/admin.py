from datetime import date, datetime, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from openpyxl import load_workbook, Workbook
from io import BytesIO
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.core.deps import require_roles, log_action
from app.core.pdf_utils import build_absence_report_pdf
from app.core.security import hash_password
from app.database import get_db
from app.models.models import (
    Student, Absence, Classroom, MealConsumption, DailyMeal, InventoryItem, InventoryMovement, AuditLog, User, Role, Schedule,
    GradeEntry, StudentProfile, TeacherProfile, SchoolScheduleConfig, Facility, SubjectWeeklyQuota,
)
from app.schemas.schemas import DashboardStats, AuditLogOut, ClassroomCreate, StudentCreate, UserCreate, GradeSubmitRequest, StudentProfileUpdate, TeacherProfileUpdate

router = APIRouter(prefix="/api/admin", tags=["الإدارة"])

@router.get("/schedule-config")
def get_schedule_config(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    row = db.query(SchoolScheduleConfig).order_by(SchoolScheduleConfig.id.desc()).first()
    if not row:
        row = SchoolScheduleConfig(); db.add(row); db.commit(); db.refresh(row)
    return {k: getattr(row, k) for k in ("id", "school_year", "mode", "morning_start", "morning_end", "afternoon_start", "afternoon_end", "lesson_minutes", "break_minutes", "working_days")}

@router.put("/schedule-config")
def update_schedule_config(payload: dict, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    row = db.query(SchoolScheduleConfig).order_by(SchoolScheduleConfig.id.desc()).first() or SchoolScheduleConfig()
    db.add(row)
    for key in ("school_year", "mode", "morning_start", "morning_end", "afternoon_start", "afternoon_end", "lesson_minutes", "break_minutes", "working_days"):
        if key in payload: setattr(row, key, payload[key])
    db.commit(); db.refresh(row)
    return {"message": "تم حفظ إعدادات دوام المؤسسة", "id": row.id}

@router.get("/facilities")
def list_facilities(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    return [{"id": f.id, "name": f.name, "kind": f.kind, "capacity": f.capacity, "is_available": f.is_available} for f in db.query(Facility).order_by(Facility.name).all()]

@router.post("/facilities")
def create_facility(payload: dict, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    row = Facility(name=str(payload.get("name", "")).strip(), kind=payload.get("kind", "CLASSROOM"), capacity=payload.get("capacity"), is_available=payload.get("is_available", True))
    if not row.name: raise HTTPException(400, "اسم المرفق مطلوب")
    db.add(row); db.commit(); db.refresh(row); return {"id": row.id, "message": "تمت إضافة المرفق"}

@router.get("/subject-quotas")
def list_subject_quotas(school_year: str = "2026/2027", current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    return [{"id": q.id, "school_year": q.school_year, "level": q.level, "subject": q.subject, "lessons_per_week": q.lessons_per_week, "minutes_per_lesson": q.minutes_per_lesson, "required": q.required} for q in db.query(SubjectWeeklyQuota).filter(SubjectWeeklyQuota.school_year == school_year).order_by(SubjectWeeklyQuota.level, SubjectWeeklyQuota.subject).all()]

@router.post("/subject-quotas")
def create_subject_quota(payload: dict, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    required = all(payload.get(k) not in (None, "") for k in ("level", "subject", "lessons_per_week"))
    if not required: raise HTTPException(400, "المستوى والمادة وعدد الحصص مطلوبة")
    row = SubjectWeeklyQuota(school_year=payload.get("school_year", "2026/2027"), level=payload["level"], subject=payload["subject"], lessons_per_week=float(payload["lessons_per_week"]), minutes_per_lesson=int(payload.get("minutes_per_lesson", 45)), required=payload.get("required", True))
    db.add(row); db.commit(); db.refresh(row); return {"id": row.id, "message": "تم حفظ الحجم الساعي للمادة"}

@router.get("/schedules")
def list_schedules(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    rows = db.query(Schedule).join(User, Schedule.teacher_id == User.id).join(Classroom, Schedule.classroom_id == Classroom.id).all()
    return [{"id": r.id, "teacher_id": r.teacher_id, "teacher": r.teacher.full_name, "subject": getattr(r, "subject", ""), "classroom_id": r.classroom_id, "classroom": r.classroom.name, "day": r.day_of_week, "period": r.period.value, "status": r.status} for r in rows]

@router.post("/schedules")
def create_schedule(payload: dict, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    teacher = db.query(User).filter(User.id == payload.get("teacher_id"), User.role == Role.TEACHER).first()
    classroom = db.query(Classroom).filter(Classroom.id == payload.get("classroom_id")).first()
    if not teacher or not classroom: raise HTTPException(400, "الأستاذ أو القسم غير موجود")
    day, period = int(payload.get("day")), payload.get("period")
    if db.query(Schedule).filter(Schedule.teacher_id == teacher.id, Schedule.day_of_week == day, Schedule.period == period).first(): raise HTTPException(409, "الأستاذ لديه حصة أخرى في نفس التوقيت")
    if db.query(Schedule).filter(Schedule.classroom_id == classroom.id, Schedule.day_of_week == day, Schedule.period == period).first(): raise HTTPException(409, "القسم مرتبط بحصة أخرى في نفس التوقيت")
    row = Schedule(teacher_id=teacher.id, classroom_id=classroom.id, day_of_week=day, period=period)
    db.add(row); db.commit(); db.refresh(row)
    return {"id": row.id, "message": "تمت إضافة الحصة"}

@router.delete("/schedules/{schedule_id}")
def delete_schedule(schedule_id: int, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    row = db.query(Schedule).filter(Schedule.id == schedule_id).first()
    if not row: raise HTTPException(404, "الحصة غير موجودة")
    db.delete(row); db.commit(); return {"message": "تم حذف الحصة"}

@router.post("/schedules/approve")
def approve_schedules(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    rows = db.query(Schedule).filter(Schedule.status != "APPROVED").all()
    for row in rows: row.status = "APPROVED"
    db.commit(); log_action(db, current_user.id, "APPROVE_SCHEDULE", "Schedule", None, {"count": len(rows)})
    return {"status": "APPROVED", "count": len(rows)}

@router.get("/schedules/export.xlsx")
def export_schedules_xlsx(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    rows = db.query(Schedule).join(User, Schedule.teacher_id == User.id).join(Classroom, Schedule.classroom_id == Classroom.id).order_by(Schedule.day_of_week, Schedule.period).all()
    book = Workbook(); sheet = book.active; sheet.title = "التوزيع الأسبوعي"
    sheet.append(["الجمهورية الجزائرية الديمقراطية الشعبية"]); sheet.append(["التوزيع الأسبوعي لعمل الأساتذة"]); sheet.append(["اليوم", "الفترة", "الأستاذ", "المادة", "القسم"])
    days = ["الأحد", "الإثنين", "الثلاثاء", "الأربعاء", "الخميس"]
    for r in rows: sheet.append([days[r.day_of_week] if r.day_of_week < 5 else r.day_of_week, "صباحية" if r.period.value == "MORNING" else "مسائية", r.teacher.full_name, r.subject or "—", r.classroom.name])
    for col in "ABCDE": sheet.column_dimensions[col].width = 24
    output = BytesIO(); book.save(output); output.seek(0); log_action(db, current_user.id, "EXPORT_SCHEDULE", "Schedule", None, {"count": len(rows)})
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="school-weekly-schedule.xlsx"'})

@router.get("/schedules/coverage")
def schedule_coverage(school_year: str = "2026/2027", current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    quotas = db.query(SubjectWeeklyQuota).filter(SubjectWeeklyQuota.school_year == school_year).all()
    result = []
    for q in quotas:
        assigned = db.query(Schedule).filter(Schedule.subject == q.subject).count()
        expected = int(q.lessons_per_week) * db.query(Classroom).count()
        result.append({"level": q.level, "subject": q.subject, "expected": expected, "assigned": assigned, "covered": min(100, round(assigned / expected * 100)) if expected else 0})
    return result

@router.get("/teachers/workload")
def teacher_workload(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    profiles = {p.user_id: p for p in db.query(TeacherProfile).all()}
    rows = []
    for teacher in db.query(User).filter(User.role == Role.TEACHER).all():
        profile = profiles.get(teacher.id); lessons = db.query(Schedule).filter(Schedule.teacher_id == teacher.id).all()
        used = sum(next((q.minutes_per_lesson for q in db.query(SubjectWeeklyQuota).filter(SubjectWeeklyQuota.subject == s.subject).all()), 45) for s in lessons)
        target = profile.weekly_target_minutes if profile else 1440
        rows.append({"teacher_id": teacher.id, "teacher": teacher.full_name, "service_type": profile.service_type if profile else "FULL_TIME", "target_minutes": target, "assigned_minutes": used, "percentage": round(used / target * 100) if target else 0})
    return rows

@router.post("/schedules/generate")
def generate_schedules(payload: dict = {}, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    """ينشئ مسودة توزيع من الحصص المطلوبة، مع فحص تعارض الأستاذ والقسم."""
    year = payload.get("school_year", "2026/2027")
    quotas = db.query(SubjectWeeklyQuota).filter(SubjectWeeklyQuota.school_year == year).all()
    teachers = db.query(User).filter(User.role == Role.TEACHER, User.is_active == True).all()
    profiles = {p.user_id: p for p in db.query(TeacherProfile).all()}
    loads = {t.id: 0 for t in teachers}
    targets = {t.id: (profiles.get(t.id).weekly_target_minutes if profiles.get(t.id) else 1440) for t in teachers}
    classrooms = db.query(Classroom).all()
    if not quotas or not teachers or not classrooms: raise HTTPException(400, "أضف الأساتذة والأقسام والحجم الساعي أولًا")
    # إعادة التوليد تعيد بناء المسودات فقط، وتحافظ على ما اعتمده المدير.
    db.query(Schedule).filter(Schedule.status == "DRAFT").delete(synchronize_session=False)
    db.commit()
    created, skipped = 0, []
    for quota in quotas:
        matching = [t for t in teachers if quota.subject in ((profiles.get(t.id).specialty if profiles.get(t.id) else "") or "")]
        pool = matching or teachers
        teacher = min(pool, key=lambda t: loads[t.id])
        target_classrooms = [c for c in classrooms if quota.level.strip() in c.name] or classrooms
        for classroom in target_classrooms:
            for n in range(int(quota.lessons_per_week)):
                placed = False
                if loads[teacher.id] + quota.minutes_per_lesson > targets[teacher.id]:
                    skipped.append({"level": quota.level, "subject": quota.subject, "classroom": classroom.name, "reason": "تجاوز النصاب الأسبوعي للأستاذ"})
                    continue
                for day in range(5):
                    for period in ("MORNING", "AFTERNOON"):
                        if db.query(Schedule).filter(Schedule.teacher_id == teacher.id, Schedule.day_of_week == day, Schedule.period == period).first(): continue
                        if db.query(Schedule).filter(Schedule.classroom_id == classroom.id, Schedule.day_of_week == day, Schedule.period == period).first(): continue
                        db.add(Schedule(teacher_id=teacher.id, classroom_id=classroom.id, day_of_week=day, period=period, subject=quota.subject))
                        loads[teacher.id] += 1; created += 1; placed = True; break
                    if placed: break
                if not placed: skipped.append({"level": quota.level, "subject": quota.subject, "classroom": classroom.name})
    db.commit()
    return {"created": created, "skipped": skipped, "message": f"تم إنشاء {created} حصة كمسودة"}


@router.get("/classrooms")
def list_classrooms(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    return [{"id": c.id, "name": c.name, "students": [{"id": s.id, "full_name": s.full_name, "barcode": s.barcode, "parent_phone": s.parent_phone} for s in c.students]} for c in db.query(Classroom).order_by(Classroom.name).all()]


@router.post("/inventory")
def create_inventory(payload: dict, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    item = InventoryItem(name=str(payload.get("name", "")).strip(), quantity=float(payload.get("quantity", 0)), unit=str(payload.get("unit", "كغ")), low_stock_threshold=float(payload.get("low_stock_threshold", 5)), image_url=payload.get("image_url"))
    if not item.name: raise HTTPException(400, "اسم المنتج مطلوب")
    db.add(item); db.commit(); db.refresh(item); log_action(db, current_user.id, "CREATE_INVENTORY", "InventoryItem", item.id)
    return {"id": item.id, "name": item.name, "quantity": item.quantity, "unit": item.unit, "image_url": item.image_url}


@router.post("/classrooms")
def create_classroom(payload: ClassroomCreate, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    classroom = Classroom(name=payload.name.strip())
    db.add(classroom); db.commit(); db.refresh(classroom)
    log_action(db, current_user.id, "CREATE_CLASSROOM", "Classroom", classroom.id)
    return {"id": classroom.id, "name": classroom.name, "students": []}


@router.post("/students")
def create_student(payload: StudentCreate, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    if not db.query(Classroom).filter(Classroom.id == payload.classroom_id).first(): raise HTTPException(404, "القسم غير موجود")
    if db.query(Student).filter(Student.barcode == payload.barcode).first(): raise HTTPException(400, "رمز التلميذ موجود مسبقًا")
    student = Student(**payload.model_dump()); db.add(student); db.commit(); db.refresh(student)
    log_action(db, current_user.id, "CREATE_STUDENT", "Student", student.id)
    return {"id": student.id, "full_name": student.full_name, "barcode": student.barcode, "parent_phone": student.parent_phone}


@router.post("/students/import")
async def import_students(file: UploadFile = File(...), classroom_id: int = Form(...), current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    """استيراد قائمة Excel. يدعم: الاسم واللقب/الاسم الكامل، ورقم التعريف أو barcode، ورقم الولي."""
    if not file.filename.lower().endswith((".xlsx", ".xlsm")): raise HTTPException(400, "يرجى رفع ملف Excel بصيغة xlsx")
    if not db.query(Classroom).filter(Classroom.id == classroom_id).first(): raise HTTPException(404, "القسم غير موجود")
    try:
        book = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True)
        sheet = book.active
        rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:
        raise HTTPException(400, f"تعذر قراءة ملف Excel: {exc}")
    def norm(v): return " ".join(str(v or "").strip().lower().replace("ـ", "").split())
    header_idx = next((i for i, row in enumerate(rows[:20]) if any(any(k in norm(cell) for k in ("nom", "اسم", "اللقب", "الاسم")) for cell in row)), None)
    if header_idx is None: raise HTTPException(400, "لم يتم العثور على صف العناوين. استخدم أعمدة الاسم واللقب ورقم التعريف")
    headers = [norm(v) for v in rows[header_idx]]
    def col(keys): return next((i for i, h in enumerate(headers) if any(k in h for k in keys)), None)
    name_col, surname_col = col(("الاسم الكامل", "الاسم", "prenom", "first")), col(("اللقب", "nom", "surname"))
    code_col = col(("barcode", "bar code", "رقم التعريف", "matricule", "الرقم")); phone_col = col(("هاتف", "phone", "ولي"))
    if name_col is None and surname_col is None: raise HTTPException(400, "عمود الاسم غير موجود")
    added, duplicate, invalid = 0, 0, []
    for line, row in enumerate(rows[header_idx + 1:], header_idx + 2):
        name = " ".join(str(row[i]).strip() for i in (surname_col, name_col) if i is not None and i < len(row) and row[i])
        code = str(row[code_col]).strip() if code_col is not None and code_col < len(row) and row[code_col] else ""
        if not name and not code: continue
        if not name or not code: invalid.append({"row": line, "reason": "الاسم أو رقم التعريف ناقص"}); continue
        if db.query(Student).filter(Student.barcode == code).first(): duplicate += 1; continue
        db.add(Student(full_name=name, barcode=code, classroom_id=classroom_id, parent_phone=str(row[phone_col]).strip() if phone_col is not None and phone_col < len(row) and row[phone_col] else None)); added += 1
    db.commit(); log_action(db, current_user.id, "IMPORT_STUDENTS", "Classroom", classroom_id, {"file": file.filename, "added": added, "duplicates": duplicate, "invalid": len(invalid)})
    return {"filename": file.filename, "added": added, "duplicates": duplicate, "invalid": invalid, "total_rows": added + duplicate + len(invalid)}


@router.post("/students/import/preview")
async def preview_student_import(file: UploadFile = File(...), current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN"))):
    """معاينة آمنة لأول صفوف Excel دون الكتابة في قاعدة البيانات."""
    if not file.filename.lower().endswith((".xlsx", ".xlsm")): raise HTTPException(400, "يرجى رفع ملف Excel بصيغة xlsx")
    try:
        book = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True); rows = list(book.active.iter_rows(values_only=True))
    except Exception as exc: raise HTTPException(400, f"تعذر قراءة ملف Excel: {exc}")
    header_idx = next((i for i, row in enumerate(rows[:20]) if any(any(k in str(cell or "").lower() for k in ("nom", "اسم", "اللقب", "الاسم")) for cell in row)), None)
    if header_idx is None: raise HTTPException(400, "لم يتم العثور على صف العناوين")
    headers = [str(v or "").strip() for v in rows[header_idx]]
    return {"header_row": header_idx + 1, "headers": headers, "rows": [[str(v or "") for v in row] for row in rows[header_idx + 1:header_idx + 6]], "total_rows": max(0, len(rows) - header_idx - 1)}


@router.get("/users")
def list_users(current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    return [{"id": u.id, "username": u.username, "full_name": u.full_name, "role": u.role.value, "active": u.is_active} for u in db.query(User).order_by(User.created_at.desc()).all()]


@router.get("/students/{student_id}/profile")
def student_profile(student_id: int, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student: raise HTTPException(404, "التلميذ غير موجود")
    p = student.profile
    return {"student": {"id": student.id, "name": student.full_name, "barcode": student.barcode, "classroom": student.classroom.name}, "profile": {k: getattr(p, k) for k in ("birth_date", "national_id", "parent_name", "parent_phone", "address", "health_notes", "photo_url")} if p else {}, "absences": [{"date": a.date.isoformat(), "period": a.period.value} for a in sorted(student.absences, key=lambda x: x.date, reverse=True)[:30]], "grades": [{"term": g.term, "subject": g.subject, "continuous": g.continuous, "exam": g.exam, "status": g.status} for g in db.query(GradeEntry).filter(GradeEntry.student_id == student_id).order_by(GradeEntry.term, GradeEntry.subject).all()], "meals": [{"date": m.meal.date.isoformat(), "period": m.meal.period.value, "description": m.meal.description} for m in student.meals[-30:]]}


@router.put("/students/{student_id}/profile")
def update_student_profile(student_id: int, payload: StudentProfileUpdate, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student: raise HTTPException(404, "التلميذ غير موجود")
    p = student.profile or StudentProfile(student_id=student_id); db.add(p)
    for key, value in payload.model_dump().items(): setattr(p, key, value or None)
    db.commit(); return {"message": "تم تحديث بطاقة التلميذ"}


@router.get("/users/{user_id}/profile")
def teacher_profile(user_id: int, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(404, "المستخدم غير موجود")
    p = db.query(TeacherProfile).filter(TeacherProfile.user_id == user_id).first()
    return {"user": {"id": user.id, "name": user.full_name, "username": user.username, "role": user.role.value}, "profile": {k: getattr(p, k) for k in ("specialty", "phone", "email", "hire_date", "address", "photo_url", "service_type", "weekly_target_minutes")} if p else {}}


@router.put("/users/{user_id}/profile")
def update_teacher_profile(user_id: int, payload: TeacherProfileUpdate, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(404, "المستخدم غير موجود")
    p = db.query(TeacherProfile).filter(TeacherProfile.user_id == user_id).first() or TeacherProfile(user_id=user_id); db.add(p)
    for key, value in payload.model_dump().items(): setattr(p, key, value or None)
    db.commit(); return {"message": "تم تحديث البطاقة المهنية"}


@router.get("/grades")
def admin_grades(term: str, subject: str, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    rows = db.query(GradeEntry).filter(GradeEntry.term == term, GradeEntry.subject == subject).order_by(GradeEntry.student_id).all()
    return [{"id": r.id, "student_id": r.student_id, "student_name": r.student.full_name, "classroom": r.student.classroom.name, "continuous": r.continuous, "exam": r.exam, "observation": r.observation, "status": r.status} for r in rows]


@router.post("/grades/approve")
def approve_grades(payload: GradeSubmitRequest, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    count = db.query(GradeEntry).filter(GradeEntry.term == payload.term, GradeEntry.subject == payload.subject, GradeEntry.status == "SUBMITTED").update({"status": "APPROVED"}, synchronize_session=False)
    db.commit(); log_action(db, current_user.id, "APPROVE_GRADES", "GradeBook", None, {"term": payload.term, "subject": payload.subject, "count": count}); return {"count": count, "status": "APPROVED"}


@router.get("/grades/export.xlsx")
def export_grades_xlsx(term: str, subject: str, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    """تصدير دفتر النقاط المعتمد بصيغة قريبة من نموذج حجز النقاط الجزائري."""
    rows = db.query(GradeEntry).filter(GradeEntry.term == term, GradeEntry.subject == subject, GradeEntry.status == "APPROVED").order_by(GradeEntry.student_id).all()
    book = Workbook()
    sheet = book.active; sheet.title = subject[:31]
    sheet.append(["الجمهورية الجزائرية الديمقراطية الشعبية"]); sheet.append(["وزارة التربية الوطنية"]); sheet.append(["وثيقة حجز النقاط"]); sheet.append([f"الفصل: {term}    المادة: {subject}"]); sheet.append([])
    sheet.append(["رقم التعريف", "اللقب والاسم", "القسم", "معدل التقويم المستمر /10", "علامة الاختبار /10", "الملاحظات"])
    for r in rows: sheet.append([r.student.barcode, r.student.full_name, r.student.classroom.name, r.continuous, r.exam, r.observation or ""])
    for col in sheet.columns:
        letter = col[0].column_letter; sheet.column_dimensions[letter].width = 24 if letter in ("A", "B") else 18
    output = BytesIO(); book.save(output); output.seek(0)
    log_action(db, current_user.id, "EXPORT_GRADES", "GradeBook", None, {"term": term, "subject": subject, "count": len(rows)})
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="school-grades.xlsx"'})


@router.get("/dashboard/stats", response_model=DashboardStats)
def dashboard_stats(
    current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")),
    db: Session = Depends(get_db),
):
    today = date.today()
    total_students = db.query(Student).count()
    absent_today = db.query(Absence).filter(Absence.date == today).count()
    meals_today = (
        db.query(MealConsumption)
        .join(DailyMeal)
        .filter(DailyMeal.date == today)
        .count()
    )
    low_stock = len([i for i in db.query(InventoryItem).all() if i.quantity <= i.low_stock_threshold])

    return DashboardStats(
        total_students=total_students,
        absent_today=absent_today,
        meals_served_today=meals_today,
        low_stock_items=low_stock,
    )


@router.get("/dashboard/overview")
def dashboard_overview(
    current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")),
    db: Session = Depends(get_db),
):
    """بيانات مركّزة للوحة المدير: الأقسام، الغياب، المخزون، والحسابات."""
    today = date.today()
    classrooms = db.query(Classroom).all()
    absence_counts = dict(
        db.query(Student.classroom_id, func.count(Absence.id))
        .join(Absence, Absence.student_id == Student.id)
        .filter(Absence.date == today)
        .group_by(Student.classroom_id)
        .all()
    )
    return {
        "date": today.isoformat(),
        "classrooms": [
            {"id": c.id, "name": c.name, "students": len(c.students),
             "absent": absence_counts.get(c.id, 0),
             "studentsList": [{"id": s.id, "full_name": s.full_name, "barcode": s.barcode} for s in c.students]}
            for c in classrooms
        ],
        "inventory": [
            {"id": i.id, "name": i.name, "quantity": i.quantity, "unit": i.unit,
             "threshold": i.low_stock_threshold,
             "low": i.quantity <= i.low_stock_threshold, "image_url": i.image_url}
            for i in db.query(InventoryItem).order_by(InventoryItem.name).all()
        ],
        "users": [
            {"id": u.id, "username": u.username, "full_name": u.full_name,
             "role": u.role.value, "active": u.is_active,
             "online": bool(u.last_login_at and (datetime.utcnow() - u.last_login_at).total_seconds() <= 90),
             "last_seen": u.last_login_at.isoformat() if u.last_login_at else None}
            for u in db.query(User).order_by(User.created_at.desc()).all()
        ],
    }


@router.get("/absences")
def list_absences(start: date | None = None, end: date | None = None, classroom_id: int | None = None, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    end = end or date.today(); start = start or end
    query = db.query(Absence).join(Student).filter(Absence.date >= start, Absence.date <= end)
    if classroom_id: query = query.filter(Student.classroom_id == classroom_id)
    return [{"id": a.id, "date": a.date.isoformat(), "period": a.period.value, "student_id": a.student_id, "student_name": a.student.full_name, "classroom": a.student.classroom.name} for a in query.order_by(Absence.date.desc(), Classroom.name, Student.full_name).all()]


@router.get("/audit-logs", response_model=list[AuditLogOut])
def audit_logs(
    current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")),
    db: Session = Depends(get_db),
    limit: int = 100,
):
    return db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(limit).all()


@router.get("/reports/absences.xlsx")
def absences_report_xlsx(start: date | None = None, end: date | None = None, current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")), db: Session = Depends(get_db)):
    end = end or date.today(); start = start or (end - timedelta(days=7))
    rows = db.query(Absence).join(Student).filter(Absence.date >= start, Absence.date <= end).order_by(Absence.date, Student.full_name).all()
    book = Workbook(); sheet = book.active; sheet.title = "سجل الغياب"
    sheet.append(["الجمهورية الجزائرية الديمقراطية الشعبية"]); sheet.append(["سجل الغياب", start.isoformat(), end.isoformat()]); sheet.append(["التاريخ", "التلميذ", "القسم", "الفترة", "الباركود"])
    for a in rows: sheet.append([a.date.isoformat(), a.student.full_name, a.student.classroom.name, "صباحية" if a.period.value == "MORNING" else "مسائية", a.student.barcode])
    output = BytesIO(); book.save(output); output.seek(0); log_action(db, current_user.id, "EXPORT_ABSENCES", "Absence", None, {"count": len(rows)})
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="school-absences.xlsx"'})

@router.get("/reports/absences.pdf")
def absences_report_pdf(
    start: date | None = None,
    end: date | None = None,
    current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")),
    db: Session = Depends(get_db),
):
    """تقرير غياب PDF لفترة زمنية (افتراضيًا آخر 7 أيام). مثال: ?start=2026-09-01&end=2026-09-07"""
    end = end or date.today()
    start = start or (end - timedelta(days=7))
    if start > end:
        raise HTTPException(status_code=400, detail="تاريخ البداية يجب أن يسبق تاريخ النهاية")

    records = (
        db.query(Absence)
        .join(Student, Absence.student_id == Student.id)
        .join(Classroom, Student.classroom_id == Classroom.id)
        .filter(Absence.date >= start, Absence.date <= end)
        .order_by(Absence.date, Classroom.name)
        .all()
    )
    rows = [
        {
            "date": a.date,
            "classroom": a.student.classroom.name,
            "student": a.student.full_name,
            "period": a.period.value,
        }
        for a in records
    ]

    buffer = BytesIO()
    build_absence_report_pdf(buffer, start, end, rows)
    buffer.seek(0)

    log_action(
        db, current_user.id, "EXPORT_REPORT", "AbsenceReport", None,
        {"start": start.isoformat(), "end": end.isoformat(), "rows": len(rows)},
    )

    filename = f"absence-report-{start.isoformat()}_to_{end.isoformat()}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/users")
def create_user(
    username: str,
    password: str,
    full_name: str,
    role: str,
    current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")),
    db: Session = Depends(get_db),
):
    if db.query(User).filter(User.username == username).first():
        raise HTTPException(status_code=400, detail="اسم المستخدم موجود مسبقًا")
    try:
        role_enum = Role(role)
    except ValueError:
        raise HTTPException(status_code=400, detail="دور غير صحيح")

    user = User(username=username, password_hash=hash_password(password), full_name=full_name, role=role_enum)
    db.add(user)
    db.commit()
    db.refresh(user)
    log_action(db, current_user.id, "CREATE_USER", "User", user.id)
    return {"id": user.id, "username": user.username, "role": user.role.value}


@router.put("/users/{user_id}/deactivate")
def deactivate_user(
    user_id: int,
    current_user: User = Depends(require_roles("ADMIN", "SUPER_ADMIN")),
    db: Session = Depends(get_db),
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    user.is_active = False
    db.commit()
    log_action(db, current_user.id, "DEACTIVATE_USER", "User", user_id)
    return {"message": "تم تعطيل الحساب"}
